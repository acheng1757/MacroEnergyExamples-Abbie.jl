from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import io

SRC_PATH = '/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/sc_esc_lcoe/LCOE_SYNTHETIC_Ethylene.xlsx'
OUT_PATH = '/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/LCOE_Ethylene_Scenarios.xlsx'

wb_src = load_workbook(SRC_PATH, data_only=True)
ws_src = wb_src.active

# ── Source column indices ──────────────────────────────────────────────────────
COL_ID        = 39
COL_COMMODITY = 40
AU_BC         = list(range(47, 56))
AU_BC_HEADERS = [ws_src.cell(row=3, column=c).value for c in AU_BC]

SCENARIOS = [
    ("CHEM SEQUESTRATION - CARBON INTENSITY",    41, 56, 57, 58),
    ("CHEM SEQUESTRATION - CREDIT",              42, 59, 60, 61),
    ("NO CHEM SEQUESTRATION - CARBON INTENSITY", 44, 62, 63, 64),
    ("NO CHEM SEQUESTRATION - CREDIT",           45, 65, 66, 67),
]
# Each tuple: (scenario_name, lcoe_col, ci_col, capture_col, emissions_col)

DATA_START = 4
DATA_END   = 66

# ── Color palette ─────────────────────────────────────────────────────────────
SCENARIO_COLORS = ["1F4E79", "2E75B6", "1D6F42", "375623"]
SCENARIO_LIGHT  = ["D6E4F0", "DDEEFF", "D6EAD6", "E2EFDA"]

COST_COLORS = [
    "#4472C4",  # h2_consumption
    "#70AD47",  # h2_production
    "#ED7D31",  # elec_consumption
    "#A9D18E",  # ethane_consumption
    "#FF0000",  # natgas_consumption
    "#9DC3E6",  # natgas_production
    "#7030A0",  # investment_cost
    "#C55A11",  # fixed_om_cost
    "#BF8F00",  # variable_om_cost
    "#833C00",  # modified capture cost
    "#375623",  # modified emissions cost
]

OUTPUT_COLS = (["id", "commodity"] + AU_BC_HEADERS +
               ["CI", "modified capture cost", "modified emissions cost", "LCOE ($/t-ethylene)"])
COST_COLS   = AU_BC_HEADERS + ["modified capture cost", "modified emissions cost"]

THIN   = Side(style='thin', color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ── Read and sort scenario data ───────────────────────────────────────────────
def read_scenario_data(lcoe_col, ci_col, cap_col, emis_col):
    rows = []
    for src_row in range(DATA_START, DATA_END + 1):
        id_val = ws_src.cell(row=src_row, column=COL_ID).value
        if id_val is None:
            continue
        row = {
            "id":        id_val,
            "commodity": ws_src.cell(row=src_row, column=COL_COMMODITY).value,
        }
        for i, c in enumerate(AU_BC):
            row[AU_BC_HEADERS[i]] = ws_src.cell(row=src_row, column=c).value or 0
        row["CI"]                      = ws_src.cell(row=src_row, column=ci_col).value or 0
        row["modified capture cost"]   = ws_src.cell(row=src_row, column=cap_col).value or 0
        row["modified emissions cost"] = ws_src.cell(row=src_row, column=emis_col).value or 0
        row["LCOE ($/t-ethylene)"]     = ws_src.cell(row=src_row, column=lcoe_col).value or 0
        rows.append(row)
    rows.sort(key=lambda r: r["LCOE ($/t-ethylene)"], reverse=True)
    return rows

# ── Build stacked bar chart ───────────────────────────────────────────────────
def make_chart(rows, scenario_name):
    ids = [r["id"] for r in rows]
    n   = len(ids)

    pos_vals = {c: [] for c in COST_COLS}
    neg_vals = {c: [] for c in COST_COLS}
    for r in rows:
        for c in COST_COLS:
            v = r.get(c, 0) or 0
            pos_vals[c].append(max(v, 0))
            neg_vals[c].append(min(v, 0))

    colors_map = dict(zip(COST_COLS, COST_COLORS))

    fig, ax = plt.subplots(figsize=(max(16, n * 0.35), 7))
    fig.patch.set_facecolor('#F7F9FC')
    ax.set_facecolor('#F7F9FC')

    x          = np.arange(n)
    bar_w      = 0.6
    pos_bottom = np.zeros(n)
    neg_bottom = np.zeros(n)

    for col in COST_COLS:
        pv = np.array(pos_vals[col])
        nv = np.array(neg_vals[col])
        ax.bar(x, pv, bar_w, bottom=pos_bottom, color=colors_map[col], label=col, zorder=3)
        ax.bar(x, nv, bar_w, bottom=neg_bottom, color=colors_map[col], zorder=3)
        pos_bottom += pv
        neg_bottom += nv

    lcoe_vals = [r["LCOE ($/t-ethylene)"] for r in rows]
    ax.plot(x, lcoe_vals, 'k--o', linewidth=1.5, markersize=4, label="LCOE", zorder=5)

    ax.axhline(0, color='black', linewidth=0.8, zorder=4)
    ax.set_xticks(x)
    ax.set_xticklabels(ids, rotation=45, ha='right', fontsize=8)
    ax.set_ylabel("$/t-ethylene", fontsize=10)
    ax.set_title(scenario_name, fontsize=12, fontweight='bold', pad=12)
    ax.grid(axis='y', alpha=0.3, zorder=0)
    ax.legend(loc='upper right', fontsize=7, framealpha=0.85, ncol=2, bbox_to_anchor=(1.0, 1.0))

    plt.tight_layout()
    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=150, bbox_inches='tight')
    plt.close()
    buf.seek(0)
    return buf

# ── Build workbook ────────────────────────────────────────────────────────────
new_wb = Workbook()
new_ws = new_wb.active
new_ws.title = "Scenarios"

write_row = 1

for s_idx, (scenario_name, lcoe_col, ci_col, cap_col, emis_col) in enumerate(SCENARIOS):
    rows      = read_scenario_data(lcoe_col, ci_col, cap_col, emis_col)
    sc_color  = SCENARIO_COLORS[s_idx]
    light_col = SCENARIO_LIGHT[s_idx]

    # Title row
    title_cell = new_ws.cell(row=write_row, column=1, value=scenario_name)
    title_cell.font      = Font(bold=True, size=13, color="FFFFFF")
    title_cell.fill      = PatternFill("solid", fgColor=sc_color)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    new_ws.row_dimensions[write_row].height = 22
    new_ws.merge_cells(start_row=write_row, start_column=1,
                       end_row=write_row, end_column=len(OUTPUT_COLS))
    write_row += 1

    # Header row
    for c_idx, header in enumerate(OUTPUT_COLS, start=1):
        cell            = new_ws.cell(row=write_row, column=c_idx, value=header)
        cell.font       = Font(bold=True, color="FFFFFF", size=9)
        cell.fill       = PatternFill("solid", fgColor=sc_color)
        cell.alignment  = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border     = BORDER
    new_ws.row_dimensions[write_row].height = 30
    write_row += 1

    # Data rows
    for r_idx, row_data in enumerate(rows):
        fill_color = light_col if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, col in enumerate(OUTPUT_COLS, start=1):
            val            = row_data.get(col)
            cell           = new_ws.cell(row=write_row, column=c_idx, value=val)
            cell.fill      = PatternFill("solid", fgColor=fill_color)
            cell.border    = BORDER
            cell.alignment = Alignment(
                horizontal="center" if c_idx > 2 else "left",
                vertical="center"
            )
            if isinstance(val, float):
                cell.number_format = '#,##0.00'
        write_row += 1

    write_row += 3  # 2 blank rows + 1

# Column widths
for col_idx, col_name in enumerate(OUTPUT_COLS, start=1):
    col_letter = get_column_letter(col_idx)
    new_ws.column_dimensions[col_letter].width = min(len(str(col_name)) + 3, 22)
new_ws.column_dimensions['A'].width = 24
new_ws.column_dimensions['B'].width = 12

# ── Charts sheet ──────────────────────────────────────────────────────────────
chart_ws = new_wb.create_sheet("Charts")
chart_ws.sheet_view.showGridLines = False

chart_row = 1
for scenario_name, lcoe_col, ci_col, cap_col, emis_col in SCENARIOS:
    rows    = read_scenario_data(lcoe_col, ci_col, cap_col, emis_col)
    img_buf = make_chart(rows, scenario_name)
    img        = XLImage(img_buf)
    img.width  = 1100
    img.height = 480
    chart_ws.add_image(img, f"A{chart_row}")
    chart_row += 33

new_wb.save(OUT_PATH)
print(f"Done! Saved to {OUT_PATH}")
print("DONE SCENARIOS XLSX GENERATED")