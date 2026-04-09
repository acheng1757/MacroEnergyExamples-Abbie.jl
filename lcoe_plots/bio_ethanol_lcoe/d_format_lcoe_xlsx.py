import pandas as pd
import numpy as np
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
import shutil
import sys
import os

INPUT = '/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/bio_ethanol_lcoe/LCOE_BIOETHANOL.xlsx'
OUTPUT = '/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/bio_ethanol_lcoe/LCOE_BIOETHANOL.xlsx'

# ── 1. Read source data ──────────────────────────────────────────────────────
raw = pd.read_excel(INPUT, sheet_name='Sheet1', header=None)

# Row 2 (idx=2) holds the sub-headers we'll use as column names
# Rows 3–6 hold data (0-indexed)

# Excel column letter -> 0-based pandas index mapping
def xl(letter):
    result = 0
    for c in letter.upper():
        result = result * 26 + (ord(c) - ord('A') + 1)
    return result - 1

# Build a named DataFrame from the data rows (rows 3-6, 0-indexed)
data_rows = raw.iloc[3:].copy().reset_index(drop=True)

# Collect all columns we might need
all_cols = ['AK', 'AL', 'AM', 'AN', 'AR', 'AS', 'AT', 'AU', 'AV', 'AW', 'AX',
            'AY', 'AZ', 'BA', 'BB', 'BC', 'BD']

col_idx = {c: xl(c) for c in all_cols}

# Sub-header labels for display (from row idx=2)
sub_headers = {c: str(raw.iloc[2, col_idx[c]]) for c in all_cols}

# Also grab the row-1 group labels for merged header rows
group_labels = {c: str(raw.iloc[1, col_idx[c]]) if not pd.isna(raw.iloc[1, col_idx[c]]) else '' for c in all_cols}

# Build working dataframe
df = pd.DataFrame({c: data_rows.iloc[:, col_idx[c]].values for c in all_cols})
df.columns = all_cols

# Drop rows that are fully NaN
df = df.dropna(how='all').reset_index(drop=True)

# ── 2. Define table column sets ──────────────────────────────────────────────
# Shared columns: AK, AL, AR-AX  (AR=AS=AT=AU=AV=AW=AX)
shared = ['AK', 'AL', 'AR', 'AS', 'AT', 'AU', 'AV', 'AW', 'AX']
t1_cols = shared + ['AY', 'AZ', 'BA', 'AM', 'AN']   # Table 1
t2_cols = shared + ['BB', 'BC', 'BD', 'AN']           # Table 2

# Sort
t1 = df[t1_cols].sort_values(['AM', 'AN']).reset_index(drop=True)
t2 = df[t2_cols].sort_values(['AN']).reset_index(drop=True)

# ── 3. Load workbook and add new sheet ───────────────────────────────────────
if os.path.abspath(INPUT) != os.path.abspath(OUTPUT):
    shutil.copy(INPUT, OUTPUT)
wb = load_workbook(OUTPUT)

SHEET_NAME = 'Analysis'
if SHEET_NAME in wb.sheetnames:
    del wb[SHEET_NAME]
ws = wb.create_sheet(SHEET_NAME)

# ── 4. Styling helpers ───────────────────────────────────────────────────────
HDR_FILL   = PatternFill('solid', start_color='1F4E79')   # dark blue
HDR2_FILL  = PatternFill('solid', start_color='2E75B6')   # medium blue
DATA_FILL  = PatternFill('solid', start_color='D6E4F0')   # light blue
DATA_ALT   = PatternFill('solid', start_color='EBF3FB')   # alternate row
WHITE_FILL = PatternFill('solid', start_color='FFFFFF')

HDR_FONT   = Font(name='Arial', bold=True, color='FFFFFF', size=9)
HDR2_FONT  = Font(name='Arial', bold=True, color='FFFFFF', size=9)
DATA_FONT  = Font(name='Arial', size=9)
TITLE_FONT = Font(name='Arial', bold=True, size=11)

thin = Side(style='thin', color='B0C4D8')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT   = Alignment(horizontal='left', vertical='center', wrap_text=True)

def apply_header(cell, text, fill=HDR_FILL, font=HDR_FONT):
    cell.value = text
    cell.fill = fill
    cell.font = font
    cell.alignment = CENTER
    cell.border = border

def apply_data(cell, value, row_num, num_fmt=None):
    cell.value = value
    cell.fill = DATA_FILL if row_num % 2 == 0 else DATA_ALT
    cell.font = DATA_FONT
    cell.alignment = CENTER
    cell.border = border
    if num_fmt:
        cell.number_format = num_fmt

def num_fmt_for(col):
    # LCOE columns → 2 decimal places; ID/commodity → text style
    text_cols = {'AK', 'AL'}
    if col in text_cols:
        return '@'
    return '#,##0.00'

# ── 5. Write tables ───────────────────────────────────────────────────────────
TABLE_START_ROW = 2   # 1-indexed
GAP_COLS = 1          # one empty column between tables

# Table 1
t1_start_col = 1
t1_ncols = len(t1_cols)

# Title
title1 = ws.cell(row=TABLE_START_ROW, column=t1_start_col, value='Table 1 — Sorted by AM, AN')
title1.font = TITLE_FONT
ws.merge_cells(start_row=TABLE_START_ROW, start_column=t1_start_col,
               end_row=TABLE_START_ROW, end_column=t1_start_col + t1_ncols - 1)
title1.alignment = LEFT

HDR_ROW = TABLE_START_ROW + 1
for j, col in enumerate(t1_cols):
    cell = ws.cell(row=HDR_ROW, column=t1_start_col + j)
    label = f"{col}\n{sub_headers[col]}"
    apply_header(cell, label)

for i, row in t1.iterrows():
    for j, col in enumerate(t1_cols):
        cell = ws.cell(row=HDR_ROW + 1 + i, column=t1_start_col + j)
        apply_data(cell, row[col], i, num_fmt_for(col))

t1_last_row = HDR_ROW + len(t1)

# Table 2
t2_start_col = t1_start_col + t1_ncols + GAP_COLS
t2_ncols = len(t2_cols)

title2 = ws.cell(row=TABLE_START_ROW, column=t2_start_col, value='Table 2 — Sorted by AN')
title2.font = TITLE_FONT
ws.merge_cells(start_row=TABLE_START_ROW, start_column=t2_start_col,
               end_row=TABLE_START_ROW, end_column=t2_start_col + t2_ncols - 1)
title2.alignment = LEFT

for j, col in enumerate(t2_cols):
    cell = ws.cell(row=HDR_ROW, column=t2_start_col + j)
    label = f"{col}\n{sub_headers[col]}"
    apply_header(cell, label)

for i, row in t2.iterrows():
    for j, col in enumerate(t2_cols):
        cell = ws.cell(row=HDR_ROW + 1 + i, column=t2_start_col + j)
        apply_data(cell, row[col], i, num_fmt_for(col))

t2_last_row = HDR_ROW + len(t2)

# Column widths
for j, col in enumerate(t1_cols):
    letter = get_column_letter(t1_start_col + j)
    ws.column_dimensions[letter].width = 18 if col in ('AK', 'AL') else 14
for j, col in enumerate(t2_cols):
    letter = get_column_letter(t2_start_col + j)
    ws.column_dimensions[letter].width = 18 if col in ('AK', 'AL') else 14
ws.row_dimensions[HDR_ROW].height = 42

# ── 7. Save ───────────────────────────────────────────────────────────────────
wb.save(OUTPUT)
print(f'Saved to {OUTPUT}')