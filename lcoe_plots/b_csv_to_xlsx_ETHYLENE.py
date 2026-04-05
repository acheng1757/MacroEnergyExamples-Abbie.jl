import sys
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.append("/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/")
from a_json_to_csv_ETHYLENE import json_to_csv_transforms

ASSETS_PATH = "/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/9Zone_US/all_168hr/assets/"
XLSX_PATH = "/Users/abbie/MacroEnergy-Abbie.jl/MacroEnergyExamples/lcoe_plots/LCOE_Ethylene_Macro.xlsx"

json_files = [
    "thermalsteamcracker.json",
    "thermalsteamcracker_retrofit_option.json",
]

all_rows = []
for json_file in json_files:
    with open(ASSETS_PATH + json_file) as f:
        data = json.load(f)
    rows = json_to_csv_transforms(data)
    all_rows.extend(rows)

FIELDS = [
    "id", "commodity",
    "h2_consumption", "h2_production", "elec_consumption",
    "ethylene_production", "natgas_consumption", "natgas_production",
    "capture_rate", "emission_rate",
    "investment_cost", "fixed_om_cost", "variable_om_cost"
]

combined_df = pd.DataFrame(all_rows, columns=FIELDS)

wb = load_workbook(XLSX_PATH)
ws = wb.active

HEADER_ROW = 3
DATA_START_ROW = 4
light_blue = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")

col_name_to_idx = {}
for cell in ws[HEADER_ROW]:
    if cell.value in FIELDS:
        col_name_to_idx[cell.value] = cell.column

for r_offset, row_data in enumerate(combined_df.itertuples(index=False)):
    for col_name, col_idx in col_name_to_idx.items():
        value = getattr(row_data, col_name)
        cell = ws.cell(row=DATA_START_ROW + r_offset, column=col_idx, value=value)
        cell.fill = light_blue

wb.save(XLSX_PATH)
print(f"Done! {len(combined_df)} rows written.")